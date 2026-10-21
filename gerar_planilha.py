import json
import os
import sys
import argparse
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd

def parse_date_str(date_val):
    if not date_val:
        return ""
    if isinstance(date_val, datetime):
        return date_val.strftime("%d/%m/%Y")
    return str(date_val).strip()

def gerar_planilha_institucional(jsonl_path, csv_path, xlsx_path, turma_nome="TADS 035-5MN - Acadweb"):
    if not os.path.exists(jsonl_path):
        print(f"[ERRO] Arquivo {jsonl_path} não encontrado.")
        return False

    records = []
    with open(jsonl_path, 'r', encoding='utf-8') as f:
        for line in f:
            line = line.strip()
            if line:
                try:
                    records.append(json.loads(line))
                except Exception as e:
                    print(f"[AVISO] Linha JSON inválida ignorada: {e}")

    if not records:
        print("[AVISO] Nenhum dado encontrado para exportar.")
        return False

    # Agrupa por aluno para permitir gerar planilhas individuais ou consolidada
    alunos = {}
    for r in records:
        aluno = r.get("student_name", "Aluno Não Identificado")
        if aluno not in alunos:
            alunos[aluno] = []
        alunos[aluno].append(r)

    # 1. Gerar CSV consolidado padronizado
    csv_rows = []
    for r in records:
        course_name = r.get("course_name", "")
        declared_hours = r.get("declared_hours", r.get("hours", 0))
        approved_hours = r.get("approved_hours", r.get("hours", 0))
        completion_date = parse_date_str(r.get("completion_date", ""))
        semester = r.get("semester", "")
        student = r.get("student_name", "")

        csv_rows.append({
            "ALUNO": student,
            "DESCRIÇÃO DA ATIVIDADE / TÍTULO DO CERTIFICADO": course_name,
            "HORAS DECLARADAS": declared_hours,
            "DATA": completion_date,
            "SEMESTRE": semester,
            "HORAS APROVEITADAS": approved_hours,
            "CATEGORIA": r.get("category", "")
        })

    df = pd.DataFrame(csv_rows)
    df.to_csv(csv_path, index=False, encoding='utf-8', sep=';')
    print(f"✅ CSV gerado com sucesso em: {csv_path} ({len(csv_rows)} registros)")

    # 2. Gerar arquivo Excel (.xlsx) fiel ao modelo institucional
    for aluno_nome, certs in alunos.items():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Planilha1"

        # Estilos
        font_header_title = Font(name="Calibri", size=11, bold=True)
        font_table_header = Font(name="Calibri", size=11, bold=True)
        font_data = Font(name="Calibri", size=11)
        font_total = Font(name="Calibri", size=11, bold=True)

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")

        # Cabeçalho Institucional
        ws["A1"] = "ATIVIDADES COMPLEMENTARES"
        ws["A1"].font = font_header_title

        ws["A2"] = f"Turma: {turma_nome}"
        ws["A2"].font = font_header_title

        ws["A4"] = f"Aluno: {aluno_nome}"
        ws["A4"].font = font_header_title

        # Títulos das Colunas (Linha 6)
        headers = [
            "DESCRIÇÃO DA ATIVIDADE / TÍTULO DO CERTIFICADO",
            "HORAS DECLARADAS",
            "DATA",
            "SEMESTRE",
            "HORAS APROVEITADAS"
        ]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=6, column=col_idx, value=header)
            cell.font = font_table_header
            cell.alignment = align_center

        # Inserção das Linhas de Dados (Linha 7 em diante)
        start_row = 7
        current_row = start_row
        for c in certs:
            c_name = c.get("course_name", "")
            declared = c.get("declared_hours", c.get("hours", 0))
            approved = c.get("approved_hours", c.get("hours", 0))
            c_date = parse_date_str(c.get("completion_date", ""))
            c_semester = c.get("semester", "")

            ws.cell(row=current_row, column=1, value=c_name).alignment = align_left
            ws.cell(row=current_row, column=2, value=int(declared)).alignment = align_center
            ws.cell(row=current_row, column=3, value=c_date).alignment = align_center
            ws.cell(row=current_row, column=4, value=c_semester).alignment = align_center
            ws.cell(row=current_row, column=5, value=int(approved)).alignment = align_center

            for col_idx in range(1, 6):
                ws.cell(row=current_row, column=col_idx).font = font_data

            current_row += 1

        # Linha em branco e depois Total
        end_data_row = current_row - 1
        total_row = current_row + 1

        ws.cell(row=total_row, column=1, value="TOTAL DE HORAS COMPLETADAS").font = font_total
        total_cell = ws.cell(row=total_row, column=5, value=f"=SUM(E{start_row}:E{end_data_row})")
        total_cell.font = font_total
        total_cell.alignment = align_center

        # Ajuste das larguras das colunas
        ws.column_dimensions['A'].width = 38.0
        ws.column_dimensions['B'].width = 18.0
        ws.column_dimensions['C'].width = 16.0
        ws.column_dimensions['D'].width = 14.0
        ws.column_dimensions['E'].width = 22.0

        # Nome do arquivo por aluno ou nome padrão
        if len(alunos) > 1:
            nome_sanitizado = "".join(x for x in aluno_nome if x.isalnum() or x in " _-").strip()
            aluno_xlsx_path = f"Soma_Horas_{nome_sanitizado}.xlsx"
        else:
            aluno_xlsx_path = xlsx_path

        wb.save(aluno_xlsx_path)
        print(f"✅ Excel formatado gerado em: {aluno_xlsx_path} ({len(certs)} certificados)")

    return True

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Exporta dados de certificados para CSV e XLSX institucional.")
    parser.add_argument('--input', default='dados_extraidos.jsonl', help="Caminho do arquivo JSONL de entrada")
    parser.add_argument('--output-csv', default='planilha_certificados.csv', help="Caminho do arquivo CSV de saída")
    parser.add_argument('--output-xlsx', default='Soma_Horas_Complementares.xlsx', help="Caminho do arquivo XLSX de saída")
    parser.add_argument('--turma', default='TADS 035-5MN - Acadweb', help="Nome/código da turma institucional")

    args = parser.parse_args()
    gerar_planilha_institucional(args.input, args.output_csv, args.output_xlsx, args.turma)
