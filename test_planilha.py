import os
import json
import unittest
import openpyxl
import pandas as pd
from gerar_planilha import gerar_planilha_institucional

class TestGerarPlanilha(unittest.TestCase):
    def setUp(self):
        self.test_jsonl = "test_dados.jsonl"
        self.test_csv = "test_planilha.csv"
        self.test_xlsx = "test_planilha.xlsx"

        records = [
            {
                "student_name": "Maria Silva Santos",
                "course_name": "Introdução a Ciência de Dados",
                "declared_hours": 15,
                "approved_hours": 15,
                "completion_date": "15/04/2024",
                "semester": "2024.1",
                "category": "Curso"
            },
            {
                "student_name": "Maria Silva Santos",
                "course_name": "Bootcamp Fullstack Avançado",
                "declared_hours": 80,
                "approved_hours": 20,
                "completion_date": "20/10/2024",
                "semester": "2024.2",
                "category": "Curso"
            },
            {
                "student_name": "Maria Silva Santos",
                "course_name": "Estágio em Desenvolvimento de Software",
                "declared_hours": 120,
                "approved_hours": 20,
                "completion_date": "01/03/2025",
                "semester": "2025.1",
                "category": "Estágio"
            }
        ]

        with open(self.test_jsonl, "w", encoding="utf-8") as f:
            for r in records:
                f.write(json.dumps(r) + "\n")

    def tearDown(self):
        for path in [self.test_jsonl, self.test_csv, self.test_xlsx, "Soma_Horas_Maria Silva Santos.xlsx"]:
            if os.path.exists(path):
                os.remove(path)

    def test_geracao_csv_e_xlsx(self):
        ok = gerar_planilha_institucional(self.test_jsonl, self.test_csv, self.test_xlsx, "TADS 035-5MN - Acadweb")
        self.assertTrue(ok)
        self.assertTrue(os.path.exists(self.test_csv))

        # 1. Validação do CSV
        df = pd.read_csv(self.test_csv, sep=';', encoding='utf-8')
        self.assertEqual(len(df), 3)
        self.assertIn("HORAS DECLARADAS", df.columns)
        self.assertIn("HORAS APROVEITADAS", df.columns)
        self.assertEqual(list(df["HORAS DECLARADAS"]), [15, 80, 120])
        self.assertEqual(list(df["HORAS APROVEITADAS"]), [15, 20, 20])

        # 2. Validação do Excel (.xlsx)
        target_xlsx = self.test_xlsx
        if not os.path.exists(target_xlsx):
            target_xlsx = "Soma_Horas_Maria Silva Santos.xlsx"
        self.assertTrue(os.path.exists(target_xlsx))

        wb = openpyxl.load_workbook(target_xlsx)
        ws = wb.active
        self.assertEqual(ws["A1"].value, "ATIVIDADES COMPLEMENTARES")
        self.assertIn("TADS 035-5MN", ws["A2"].value)
        self.assertIn("Maria Silva Santos", ws["A4"].value)

        # Cabeçalhos na linha 6
        self.assertEqual(ws["A6"].value, "DESCRIÇÃO DA ATIVIDADE / TÍTULO DO CERTIFICADO")
        self.assertEqual(ws["B6"].value, "HORAS DECLARADAS")
        self.assertEqual(ws["E6"].value, "HORAS APROVEITADAS")

        # Dados nas linhas 7, 8, 9
        self.assertEqual(ws["B7"].value, 15)
        self.assertEqual(ws["E7"].value, 15)
        self.assertEqual(ws["B8"].value, 80)
        self.assertEqual(ws["E8"].value, 20)
        self.assertEqual(ws["B9"].value, 120)
        self.assertEqual(ws["E9"].value, 20)

        # Total com fórmula de soma
        self.assertEqual(ws["A11"].value, "TOTAL DE HORAS COMPLETADAS")
        self.assertEqual(ws["E11"].value, "=SUM(E7:E9)")

if __name__ == "__main__":
    unittest.main()
